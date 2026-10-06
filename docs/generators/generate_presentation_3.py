import os
import subprocess
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def build_pptx():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Palette
    C_BG = RGBColor(248, 250, 252)       # #F8FAFC
    C_CARD = RGBColor(255, 255, 255)     # #FFFFFF
    C_BORDER = RGBColor(226, 232, 240)   # #E2E8F0
    C_NAVY = RGBColor(15, 23, 42)        # #0F172A
    C_SLATE = RGBColor(51, 65, 85)       # #334155
    C_MUTED = RGBColor(100, 116, 139)    # #64748B
    C_BLUE = RGBColor(37, 99, 235)       # #2563EB
    C_BLUE_BG = RGBColor(239, 246, 255)  # #EFF6FF
    C_EMERALD = RGBColor(5, 150, 105)    # #059669
    C_AMBER = RGBColor(217, 119, 6)      # #D97706

    def add_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = C_BG
        bg.line.fill.background()
        return bg

    def add_header(slide, eyebrow, title, slide_num):
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(1.1))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p0 = tf.paragraphs[0]
        p0.text = eyebrow.upper()
        p0.font.size = Pt(10)
        p0.font.bold = True
        p0.font.color.rgb = C_BLUE

        p1 = tf.add_paragraph()
        p1.text = title
        p1.font.size = Pt(22)
        p1.font.bold = True
        p1.font.color.rgb = C_NAVY

        # Footer
        fb = slide.shapes.add_textbox(Inches(0.8), Inches(7.0), Inches(11.733), Inches(0.4))
        ftf = fb.text_frame
        ftf.word_wrap = True
        ftf.margin_left = ftf.margin_top = ftf.margin_right = ftf.margin_bottom = 0
        fp = ftf.paragraphs[0]
        fp.text = f"IT Helpdesk & Asset Support Management System | Presentation-III: Live UI & Database Demo | Slide {slide_num} of 12"
        fp.font.size = Pt(9)
        fp.font.color.rgb = C_MUTED

    def make_content_slide(slide_num, eyebrow, title, cards_data):
        slide = prs.slides.add_slide(blank_layout)
        add_bg(slide)
        add_header(slide, eyebrow, title, slide_num)

        num_cards = len(cards_data)
        card_w = (11.733 - (num_cards - 1) * 0.3) / num_cards
        for idx, c in enumerate(cards_data):
            cx = 0.8 + idx * (card_w + 0.3)
            box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(cx), Inches(1.6), Inches(card_w), Inches(5.1))
            box.fill.solid()
            box.fill.fore_color.rgb = C_CARD
            box.line.color.rgb = C_BORDER

            btf = box.text_frame
            btf.word_wrap = True
            btf.margin_left = Inches(0.25)
            btf.margin_right = Inches(0.25)
            btf.margin_top = Inches(0.25)

            p = btf.paragraphs[0]
            p.text = c.get('header', '').upper()
            p.font.size = Pt(10)
            p.font.bold = True
            p.font.color.rgb = c.get('badge_color', C_BLUE)

            p = btf.add_paragraph()
            p.text = c.get('title', '')
            p.font.size = Pt(16)
            p.font.bold = True
            p.font.color.rgb = C_NAVY
            p.space_after = Pt(10)

            for bullet in c.get('bullets', []):
                p = btf.add_paragraph()
                p.text = f"• {bullet}"
                p.font.size = Pt(12)
                p.font.color.rgb = C_SLATE
                p.space_after = Pt(6)

        return slide

    def make_screenshot_slide(slide_num, eyebrow, title, img_path, bullets, img_left=True):
        slide = prs.slides.add_slide(blank_layout)
        add_bg(slide)
        add_header(slide, eyebrow, title, slide_num)

        if img_left:
            # Image card on left
            box_img = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(6.8), Inches(5.1))
            box_img.fill.solid()
            box_img.fill.fore_color.rgb = C_CARD
            box_img.line.color.rgb = C_BORDER
            if os.path.exists(img_path):
                slide.shapes.add_picture(img_path, Inches(0.95), Inches(1.75), Inches(6.5), Inches(4.8))

            # Bullets on right
            box_txt = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.9), Inches(1.6), Inches(4.633), Inches(5.1))
            box_txt.fill.solid()
            box_txt.fill.fore_color.rgb = C_CARD
            box_txt.line.color.rgb = C_BORDER
            tf = box_txt.text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.25)

            p = tf.paragraphs[0]
            p.text = "OPERATION HIGHLIGHTS"
            p.font.size = Pt(10)
            p.font.bold = True
            p.font.color.rgb = C_BLUE

            for b in bullets:
                p = tf.add_paragraph()
                p.text = f"• {b}"
                p.font.size = Pt(12)
                p.font.color.rgb = C_SLATE
                p.space_after = Pt(8)
        else:
            # Bullets on left
            box_txt = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(4.633), Inches(5.1))
            box_txt.fill.solid()
            box_txt.fill.fore_color.rgb = C_CARD
            box_txt.line.color.rgb = C_BORDER
            tf = box_txt.text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.25)

            p = tf.paragraphs[0]
            p.text = "OPERATION HIGHLIGHTS"
            p.font.size = Pt(10)
            p.font.bold = True
            p.font.color.rgb = C_BLUE

            for b in bullets:
                p = tf.add_paragraph()
                p.text = f"• {b}"
                p.font.size = Pt(12)
                p.font.color.rgb = C_SLATE
                p.space_after = Pt(8)

            # Image on right
            box_img = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.733), Inches(1.6), Inches(6.8), Inches(5.1))
            box_img.fill.solid()
            box_img.fill.fore_color.rgb = C_CARD
            box_img.line.color.rgb = C_BORDER
            if os.path.exists(img_path):
                slide.shapes.add_picture(img_path, Inches(5.883), Inches(1.75), Inches(6.5), Inches(4.8))

        return slide

    # SLIDE 1: Title Slide
    s1 = prs.slides.add_slide(blank_layout)
    add_bg(s1)

    card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(0.8), Inches(11.333), Inches(5.9))
    card.fill.solid()
    card.fill.fore_color.rgb = C_CARD
    card.line.color.rgb = C_BORDER

    tb = s1.shapes.add_textbox(Inches(1.5), Inches(1.2), Inches(10.333), Inches(3.0))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "DATABASE MANAGEMENT SYSTEMS (DBMS) COURSE PROJECT"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = C_BLUE

    p = tf.add_paragraph()
    p.text = "IT Helpdesk & Asset Support\nManagement System"
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = C_NAVY
    p.space_after = Pt(14)

    p = tf.add_paragraph()
    p.text = "Presentation-III: Live User Interface Demonstration & MySQL Relational CRUD Operations"
    p.font.size = Pt(14)
    p.font.color.rgb = C_SLATE

    # Left: Student
    ib1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(4.5), Inches(4.8), Inches(1.8))
    ib1.fill.solid()
    ib1.fill.fore_color.rgb = C_BLUE_BG
    ib1.line.color.rgb = C_BORDER
    itf1 = ib1.text_frame
    itf1.margin_left = Inches(0.2)
    itf1.margin_top = Inches(0.15)
    ip = itf1.paragraphs[0]
    ip.text = "STUDENT DETAILS"
    ip.font.size = Pt(10)
    ip.font.bold = True
    ip.font.color.rgb = C_BLUE

    ip = itf1.add_paragraph()
    ip.text = "Md Aali Rahman"
    ip.font.size = Pt(16)
    ip.font.bold = True
    ip.font.color.rgb = C_NAVY

    ip = itf1.add_paragraph()
    ip.text = "Roll Number: 25WU0102156\nSection: AIML Panthers | Serial No: 39\nSlot: Wednesday, 7 October 2026 | 9:30 AM"
    ip.font.size = Pt(11)
    ip.font.color.rgb = C_SLATE

    # Right: Faculty
    ib2 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(4.5), Inches(5.0), Inches(1.8))
    ib2.fill.solid()
    ib2.fill.fore_color.rgb = C_CARD
    ib2.line.color.rgb = C_BORDER
    itf2 = ib2.text_frame
    itf2.margin_left = Inches(0.2)
    itf2.margin_top = Inches(0.15)
    ip = itf2.paragraphs[0]
    ip.text = "ACADEMIC MENTORSHIP"
    ip.font.size = Pt(10)
    ip.font.bold = True
    ip.font.color.rgb = C_BLUE

    ip = itf2.add_paragraph()
    ip.text = "Faculty Guide: Dr. Kiranmayee Adavala"
    ip.font.size = Pt(15)
    ip.font.bold = True
    ip.font.color.rgb = C_NAVY

    ip = itf2.add_paragraph()
    ip.text = "Department: Computer Science & Engineering\nInstitution: Woxsen University, School of Technology\nLive Stack: MySQL + Python Flask + Vanilla SPA"
    ip.font.size = Pt(11)
    ip.font.color.rgb = C_SLATE

    # SLIDE 2: Problem Statement
    make_content_slide(
        2, "Problem Context", "Enterprise IT Support & Asset Tracking Dilemma",
        [
            {
                "header": "Operational Bottlenecks",
                "title": "Fragmented Channels",
                "badge_color": C_AMBER,
                "bullets": [
                    "Enterprises manage hundreds of devices and support requests via chaotic emails and spreadsheets.",
                    "Untracked tickets lead to missed SLAs and zero technician accountability.",
                    "Hardware relocations occur without updating ownership or tracking warranty expiries.",
                    "Lack of immutable audit logging leads to untracked status modifications."
                ]
            },
            {
                "header": "Relational Solution",
                "title": "Normalized Database",
                "badge_color": C_EMERALD,
                "bullets": [
                    "A centralized MySQL 8.0+ relational back-end (it_helpdesk) enforcing 3NF normalization.",
                    "Clear subtype separation into Incidents vs Service Requests.",
                    "Complete audit trail of status progressions and technician assignments.",
                    "Real-time responsive web UI connected to MySQL demonstrating live CRUD."
                ]
            }
        ]
    )

    # SLIDE 3: Objectives & Scope
    make_content_slide(
        3, "Project Goals", "Objectives & Scope of Presentation-III",
        [
            {
                "header": "Core Objectives",
                "title": "Core System Objectives",
                "badge_color": C_BLUE,
                "bullets": [
                    "Deliver a fully working User Interface connected directly to MySQL (no mock data).",
                    "Demonstrate real-time INSERT of ticket records reflected in database tables.",
                    "Demonstrate permanent DELETION of records with referential cascade.",
                    "Demonstrate multi-table VIEW of records with dynamic filtering and audit timelines."
                ]
            },
            {
                "header": "Technical Scope",
                "title": "Scope & Verified Deliverables",
                "badge_color": C_EMERALD,
                "bullets": [
                    "14 Relational tables managed with foreign key cascades and restrict constraints.",
                    "Live system health monitoring with active connection verification.",
                    "Interactive 3D database visualizer inspecting schema nodes and relationships.",
                    "SQL verification playground executing live analytical queries for faculty."
                ]
            }
        ]
    )

    # SLIDE 4: Database Architecture (3NF)
    make_content_slide(
        4, "Schema Architecture", "Database Architecture & 3NF Relational Model",
        [
            {
                "header": "Schema Structure",
                "title": "14 Normalized Tables",
                "badge_color": C_BLUE,
                "bullets": [
                    "Master Entities: departments, users, categories, priorities, support_staff, warranties.",
                    "Operational Entities: assets, tickets, incidents, service_requests, assignments.",
                    "Audit & Service Entities: status_histories, resolutions, maintenance.",
                    "Storage Engine: MySQL InnoDB with ACID transaction support and foreign keys."
                ]
            },
            {
                "header": "Integrity Constraints",
                "title": "Enforced Relational Rules",
                "badge_color": C_EMERALD,
                "bullets": [
                    "ON DELETE RESTRICT: departments -> users, categories -> tickets, priorities -> tickets.",
                    "ON DELETE CASCADE: tickets -> incidents, service_requests, assignments, status_histories, resolutions.",
                    "ON DELETE SET NULL: users -> assets, warranties -> assets.",
                    "CHECK constraints on priority levels (1-5), status enums, and positive maintenance costs."
                ]
            }
        ]
    )

    # SLIDE 5: Technology Stack
    make_content_slide(
        5, "Technology Architecture", "Full-Stack System Architecture & Decoupled Design",
        [
            {
                "header": "Back-End & Database",
                "title": "MySQL & Python REST API",
                "badge_color": C_BLUE,
                "bullets": [
                    "Database Server: MySQL 8.0+ / 9.x running locally on port 3306.",
                    "Driver: mysql-connector-python with parameterized queries (zero SQL injection).",
                    "Backend: Python Flask REST API server exposing CRUD endpoints.",
                    "Transaction Management: Explicit autocommit=False with atomic commits and rollbacks."
                ]
            },
            {
                "header": "Front-End Presentation",
                "title": "Modern Single-Page Application",
                "badge_color": C_EMERALD,
                "bullets": [
                    "Vanilla HTML5 / CSS3 / ES6+ JavaScript client using the native Fetch API.",
                    "Zero heavy frontend framework dependencies for high performance.",
                    "Three.js 3D WebGL database relationship visualizer.",
                    "Live SQL verification playground for faculty evaluation."
                ]
            }
        ]
    )

    # SLIDE 6: Dashboard Overview (With Screenshot)
    make_screenshot_slide(
        6, "UI Architecture", "Operational Command Dashboard",
        "Presentation-III/screenshots/dashboard.png",
        [
            "Real-time MySQL connectivity pill: ● MySQL Connected (it_helpdesk).",
            "Dynamic aggregate counters: Total Tickets, Open vs Resolved, Hardware Assets, Users, Maintenance Spend.",
            "Interactive 3D database visualizer mapping tables and relationships.",
            "Live quick-action shortcuts to create tickets, view queues, and run SQL queries."
        ]
    )

    # SLIDE 7: VIEW Operation (With Screenshot)
    make_screenshot_slide(
        7, "CRUD Operation 1", "VIEW: Multi-Table Relational Ticket Queue",
        "Presentation-III/screenshots/tickets-view.png",
        [
            "Executes multi-table SQL JOIN across tickets, users, departments, categories, and priorities.",
            "Dynamic filtering by Status, Category, and Subtype (Incident vs Service Request).",
            "Real-time search across ticket number, title, requester name, and category.",
            "Detail modal displays full relational graph: requester details, assigned technicians, and chronological status history."
        ]
    )

    # SLIDE 8: INSERT Operation (With Screenshot)
    make_screenshot_slide(
        8, "CRUD Operation 2", "INSERT: Form-Driven Record Creation & Subtyping",
        "Presentation-III/screenshots/ticket-insert-form.png",
        [
            "Dynamic foreign key dropdowns populated from live database tables.",
            "Automated sequential ticket tracking number generation (TKT-006).",
            "Atomic multi-table transaction inserting into tickets, subtypes (incidents/service_requests), and status_histories.",
            "Immediate UI reflection: new row appears instantly in the table without page reload."
        ]
    )

    # SLIDE 9: DELETE Operation (With Screenshot)
    make_screenshot_slide(
        9, "CRUD Operation 3", "DELETE: Permanent Removal with Foreign Key Cascades",
        "Presentation-III/screenshots/ticket-before-delete.png",
        [
            "Confirmation dialog displaying exact record ID prevents accidental deletions.",
            "Executes DELETE FROM tickets WHERE ticket_id = %s directly in MySQL.",
            "ON DELETE CASCADE automatically purges sub-entities without leaving orphan rows.",
            "ON DELETE RESTRICT safely protects users and categories with active dependencies."
        ]
    )

    # SLIDE 10: Live Persistence Verification (With Screenshot)
    make_screenshot_slide(
        10, "Integrity Verification", "Live Database Persistence & Refresh Verification",
        "Presentation-III/screenshots/database-connected.png",
        [
            "State persistence verified across browser refreshes (Cmd+R / F5).",
            "Direct terminal verification via MySQL client confirms record existence in InnoDB storage.",
            "Zero mock data, fake arrays, or localStorage used — 100% real database transactions.",
            "Demonstrates end-to-end ACID consistency."
        ]
    )

    # SLIDE 11: Testing & QA
    make_content_slide(
        11, "Quality Assurance", "Automated DBMS Test Suite Execution (TC01 - TC10)",
        [
            {
                "header": "Core CRUD Test Cases",
                "title": "CRUD & Persistence Tests",
                "badge_color": C_BLUE,
                "bullets": [
                    "TC01: Database Connection to MySQL — PASS",
                    "TC02: View Tickets Relational Query — PASS",
                    "TC03: Insert New Ticket Record — PASS",
                    "TC04: Refresh & Persistence Verification — PASS",
                    "TC05: Delete Ticket from MySQL — PASS",
                    "TC06: Refresh After Delete (Remains Absent) — PASS"
                ]
            },
            {
                "header": "Constraint Tests",
                "title": "Relational Constraints & Error Handling",
                "badge_color": C_EMERALD,
                "bullets": [
                    "TC07: Form Input Validation (Missing Required Fields) — PASS",
                    "TC08: Unique Constraint Enforcement (Duplicate Email) — PASS",
                    "TC09: Foreign Key Deletion Violation (RESTRICT) — PASS",
                    "TC10: MySQL Offline Error Handling (Safe User Message) — PASS",
                    "Overall Test Result: 10/10 Test Scenarios Passed (100% Success Rate)"
                ]
            }
        ]
    )

    # SLIDE 12: Conclusion & Future Enhancements
    make_content_slide(
        12, "Summary & Future Scope", "Conclusion & Future System Enhancements",
        [
            {
                "header": "Summary of Outcomes",
                "title": "Presentation-III Accomplishments",
                "badge_color": C_EMERALD,
                "bullets": [
                    "Successfully delivered an enterprise-grade IT Helpdesk and Asset Support Management System.",
                    "Live MySQL database connectivity verified with real-time INSERT, DELETE, and VIEW operations.",
                    "Strict 3NF normalization and foreign key integrity constraints maintained throughout.",
                    "All official presentation requirements fulfilled."
                ]
            },
            {
                "header": "Future Roadmap",
                "title": "Future Enhancements",
                "badge_color": C_BLUE,
                "bullets": [
                    "Automated SLA escalation triggers and email notification webhooks.",
                    "Multi-factor authentication (MFA) and Role-Based Access Control (RBAC).",
                    "Predictive hardware maintenance forecasting using machine learning models on servicing history."
                ]
            }
        ]
    )

    out_path = "Presentation-III/Presentation-III.pptx"
    prs.save(out_path)
    print(f"Saved {out_path}")

build_pptx()
