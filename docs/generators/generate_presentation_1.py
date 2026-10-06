import os
import subprocess
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def build_pptx():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Palette
    C_BG = RGBColor(248, 250, 252)       # #F8FAFC
    C_CARD = RGBColor(255, 255, 255)     # #FFFFFF
    C_BORDER = RGBColor(226, 232, 240)   # #E2E8F0
    C_NAVY = RGBColor(15, 23, 42)        # #0F172A
    C_SLATE = RGBColor(51, 65, 85)       # #334155
    C_MUTED = RGBColor(100, 116, 139)    # #64748B
    C_BLUE = RGBColor(37, 99, 235)       # #2563EB
    C_BLUE_BG = RGBColor(239, 246, 255)  # #EFF6FF
    C_EMERALD = RGBColor(5, 150, 105)    # #059669
    C_AMBER = RGBColor(217, 119, 6)      # #D97706

    def add_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = C_BG
        bg.line.fill.background()
        return bg

    def add_header(slide, eyebrow, title, slide_num):
        # Header banner text
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(1.1))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p0 = tf.paragraphs[0]
        p0.text = eyebrow.upper()
        p0.font.size = Pt(10)
        p0.font.bold = True
        p0.font.color.rgb = C_BLUE

        p1 = tf.add_paragraph()
        p1.text = title
        p1.font.size = Pt(22)
        p1.font.bold = True
        p1.font.color.rgb = C_NAVY

        # Footer
        fb = slide.shapes.add_textbox(Inches(0.8), Inches(7.0), Inches(11.733), Inches(0.4))
        ftf = fb.text_frame
        ftf.word_wrap = True
        ftf.margin_left = ftf.margin_top = ftf.margin_right = ftf.margin_bottom = 0
        fp = ftf.paragraphs[0]
        fp.text = f"IT Helpdesk & Asset Support Management System | Presentation-I: Problem Description | Slide {slide_num}"
        fp.font.size = Pt(9)
        fp.font.color.rgb = C_MUTED

    # -------------------------------------------------------------
    # SLIDE 1: Title Slide
    # -------------------------------------------------------------
    s1 = prs.slides.add_slide(blank_layout)
    add_bg(s1)

    # Main Card
    card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(0.8), Inches(11.333), Inches(5.9))
    card.fill.solid()
    card.fill.fore_color.rgb = C_CARD
    card.line.color.rgb = C_BORDER

    tb = s1.shapes.add_textbox(Inches(1.5), Inches(1.2), Inches(10.333), Inches(3.0))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "DATABASE MANAGEMENT SYSTEMS (DBMS) COURSE PROJECT"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = C_BLUE

    p = tf.add_paragraph()
    p.text = "IT Helpdesk & Asset Support\nManagement System"
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = C_NAVY
    p.space_after = Pt(14)

    p = tf.add_paragraph()
    p.text = "Presentation-I: Problem Description, Domain Context, Objectives & Scope Specification"
    p.font.size = Pt(14)
    p.font.color.rgb = C_SLATE

    # Info boxes
    # Left box: Student
    ib1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(4.5), Inches(4.8), Inches(1.8))
    ib1.fill.solid()
    ib1.fill.fore_color.rgb = C_BLUE_BG
    ib1.line.color.rgb = C_BORDER
    itf1 = ib1.text_frame
    itf1.margin_left = Inches(0.2)
    itf1.margin_top = Inches(0.15)
    ip = itf1.paragraphs[0]
    ip.text = "STUDENT DETAILS"
    ip.font.size = Pt(10)
    ip.font.bold = True
    ip.font.color.rgb = C_BLUE

    ip = itf1.add_paragraph()
    ip.text = "Md Aali Rahman"
    ip.font.size = Pt(16)
    ip.font.bold = True
    ip.font.color.rgb = C_NAVY

    ip = itf1.add_paragraph()
    ip.text = "Roll Number: 25WU0102156\nSection: AIML Panthers | Year: 2026\nWoxsen University, School of Technology"
    ip.font.size = Pt(11)
    ip.font.color.rgb = C_SLATE

    # Right box: Faculty & Course
    ib2 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(4.5), Inches(5.0), Inches(1.8))
    ib2.fill.solid()
    ib2.fill.fore_color.rgb = C_CARD
    ib2.line.color.rgb = C_BORDER
    itf2 = ib2.text_frame
    itf2.margin_left = Inches(0.2)
    itf2.margin_top = Inches(0.15)
    ip = itf2.paragraphs[0]
    ip.text = "ACADEMIC MENTORSHIP"
    ip.font.size = Pt(10)
    ip.font.bold = True
    ip.font.color.rgb = C_BLUE

    ip = itf2.add_paragraph()
    ip.text = "Faculty Guide: Dr. Kiranmayee Adavala"
    ip.font.size = Pt(15)
    ip.font.bold = True
    ip.font.color.rgb = C_NAVY

    ip = itf2.add_paragraph()
    ip.text = "Department: Computer Science & Engineering\nCourse Code: DBMS Course Project\nDatabase: MySQL 8.0+ / 9.x Relational Engine"
    ip.font.size = Pt(11)
    ip.font.color.rgb = C_SLATE

    # Helper function to create standard 2-card or 3-card slide
    def make_content_slide(slide_num, eyebrow, title, cards_data):
        slide = prs.slides.add_slide(blank_layout)
        add_bg(slide)
        add_header(slide, eyebrow, title, slide_num)

        num_cards = len(cards_data)
        card_w = (11.733 - (num_cards - 1) * 0.3) / num_cards
        for idx, c in enumerate(cards_data):
            cx = 0.8 + idx * (card_w + 0.3)
            box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(cx), Inches(1.6), Inches(card_w), Inches(5.1))
            box.fill.solid()
            box.fill.fore_color.rgb = C_CARD
            box.line.color.rgb = C_BORDER

            btf = box.text_frame
            btf.word_wrap = True
            btf.margin_left = Inches(0.25)
            btf.margin_right = Inches(0.25)
            btf.margin_top = Inches(0.25)

            p = btf.paragraphs[0]
            p.text = c.get('header', '').upper()
            p.font.size = Pt(10)
            p.font.bold = True
            p.font.color.rgb = c.get('badge_color', C_BLUE)

            p = btf.add_paragraph()
            p.text = c.get('title', '')
            p.font.size = Pt(16)
            p.font.bold = True
            p.font.color.rgb = C_NAVY
            p.space_after = Pt(10)

            for bullet in c.get('bullets', []):
                p = btf.add_paragraph()
                p.text = f"• {bullet}"
                p.font.size = Pt(12)
                p.font.color.rgb = C_SLATE
                p.space_after = Pt(6)

        return slide

    # SLIDE 2: Problem Statement & Context
    make_content_slide(
        2, "Industry Background", "IT Service Operations & Existing Operational Bottlenecks",
        [
            {
                "header": "Enterprise Context",
                "title": "Modern IT Infrastructure",
                "badge_color": C_BLUE,
                "bullets": [
                    "Educational institutions and enterprises deploy hundreds of diverse computing assets (workstations, laptops, network routers, monitors).",
                    "Continuous IT support requests are raised daily by faculty, administrative personnel, researchers, and students.",
                    "Operations require coordination between multi-tiered support staff, hardware inventory, and manufacturer warranties."
                ]
            },
            {
                "header": "Current Deficiencies",
                "title": "Fragmented Management",
                "badge_color": C_AMBER,
                "bullets": [
                    "Widespread reliance on unstructured emails, spreadsheets, and informal communication.",
                    "Tickets are frequently lost, delayed, or untracked across busy technical shifts.",
                    "Lack of clear technician assignment leads to duplicate efforts or unattended high-priority outages.",
                    "Hardware movements occur without updating user ownership or warranty expiries."
                ]
            }
        ]
    )

    # SLIDE 3: Problem Statement
    make_content_slide(
        3, "Core Problem", "Formal Problem Statement & Academic Motivation",
        [
            {
                "header": "Problem Formulation",
                "title": "The Core Challenge",
                "badge_color": C_BLUE,
                "bullets": [
                    "The absence of a centralized, normalized relational database for IT service operations causes severe data redundancy, inconsistent state lifecycles, and untracked financial repair costs.",
                    "Manual spreadsheets lack ACID compliance, fail to enforce foreign key referential integrity, and provide zero immutable audit trails for ticket status changes.",
                    "There is no structured differentiation between unplanned incident disruptions and routine service requests."
                ]
            },
            {
                "header": "Adverse Consequences",
                "title": "Impact on Organizations",
                "badge_color": C_AMBER,
                "bullets": [
                    "High Average Time to Resolution (MTTR) due to lost ticket history.",
                    "Financial leakage from paying out-of-pocket for equipment repairs that are covered under valid manufacturer warranties.",
                    "Zero visibility into department-wise IT resource consumption and technician workloads.",
                    "Compromised audit compliance and operational accountability."
                ]
            }
        ]
    )

    # SLIDE 4: Proposed System
    make_content_slide(
        4, "Proposed Solution", "Centralized MySQL IT Helpdesk & Asset Support System",
        [
            {
                "header": "Relational Centralization",
                "title": "Unified Relational Model",
                "badge_color": C_EMERALD,
                "bullets": [
                    "Engineered with a 14-table Third Normal Form (3NF) relational database in MySQL 8.0+ / 9.x.",
                    "Enforces strict referential integrity with ON DELETE RESTRICT, CASCADE, and SET NULL constraints.",
                    "Eliminates data redundancy across departments, users, assets, categories, and technicians."
                ]
            },
            {
                "header": "Core Architectural Pillars",
                "title": "Key Functional Subsystems",
                "badge_color": C_BLUE,
                "bullets": [
                    "Decoupled Ticketing Architecture: Clear subtype classification into Incidents vs Service Requests.",
                    "Hardware Lifecycle & Warranty Engine: Asset tagging, serial number tracking, and maintenance cost logging.",
                    "Immutable Status Audit Trail: Automatic timestamped tracking of all ticket status progressions.",
                    "Multi-tier Access & Interactive Interfaces: Both terminal CLI and responsive modern web interface."
                ]
            }
        ]
    )

    # SLIDE 5: Project Objectives
    make_content_slide(
        5, "Target Goals", "Project SMART Objectives",
        [
            {
                "header": "Objective 1 & 2",
                "title": "Ticketing & Auditing",
                "badge_color": C_BLUE,
                "bullets": [
                    "Centralize Support Queues: Standardize ticket logging across 5 severity levels and 5 operational categories.",
                    "Immutable Status Audit Logs: Record every status change (Open -> In Progress -> Resolved -> Closed) in status_histories with exact timestamps."
                ]
            },
            {
                "header": "Objective 3 & 4",
                "title": "Assets & Integrity",
                "badge_color": C_EMERALD,
                "bullets": [
                    "Asset & Warranty Linkage: Map physical equipment to custodians, warranty expiration schedules, and maintenance expenses.",
                    "Strict Relational Integrity: Enforce primary keys, unique constraints, and foreign keys across all 14 tables."
                ]
            },
            {
                "header": "Objective 5",
                "title": "Demonstration",
                "badge_color": C_AMBER,
                "bullets": [
                    "Working Demonstration: Deliver a fully functional Python interface directly connected to MySQL with verified CRUD operations."
                ]
            }
        ]
    )

    # SLIDE 6: Project Scope & Boundaries
    make_content_slide(
        6, "Boundary Definition", "Scope of the System & Operational Boundaries",
        [
            {
                "header": "In-Scope Modules",
                "title": "System Boundaries (In-Scope)",
                "badge_color": C_EMERALD,
                "bullets": [
                    "Department and user master records with unique corporate emails.",
                    "Ticket lifecycle management: categorization, prioritization, assignment.",
                    "Specialized subtype tracking for Incidents vs Service Requests.",
                    "Asset inventory with unique asset tags, serial numbers, and warranty contracts.",
                    "Maintenance expense tracking and total cost calculation per department.",
                    "Relational reporting: JOIN queries, aggregate workloads, audit timelines."
                ]
            },
            {
                "header": "Out-of-Scope Elements",
                "title": "Explicit Exclusions (Out-of-Scope)",
                "badge_color": C_MUTED,
                "bullets": [
                    "Commercial payment gateway processing (maintenance costs are accounting logs).",
                    "Automated machine learning ticket auto-triage (reserved for future work).",
                    "Direct hardware telemetric scanning or IoT board firmware flashing.",
                    "Public open internet self-registration (user directory is managed by IT admin)."
                ]
            }
        ]
    )

    # SLIDE 7: Users & Stakeholders
    make_content_slide(
        7, "Stakeholder Analysis", "Users, Roles & System Permissions",
        [
            {
                "header": "Role 1",
                "title": "End-Users (Employees & Students)",
                "badge_color": C_BLUE,
                "bullets": [
                    "Raise new support tickets with titles, descriptions, categories, and priorities.",
                    "Specify whether the request is an Incident or a Service Request.",
                    "Track real-time progress and view chronological status audit trails.",
                    "View personal allocated hardware assets and warranty status."
                ]
            },
            {
                "header": "Role 2",
                "title": "IT Support Staff (Technicians)",
                "badge_color": C_EMERALD,
                "bullets": [
                    "Inspect assigned tickets filtered by technical specialization.",
                    "Update ticket status from Open to In Progress and record work notes.",
                    "Log official resolution narratives upon successful issue remediation.",
                    "Review asset maintenance history and warranty coverage."
                ]
            },
            {
                "header": "Role 3",
                "title": "IT Helpdesk Administrators",
                "badge_color": C_AMBER,
                "bullets": [
                    "Oversee queue metrics, active tickets, and technician workload balances.",
                    "Register new hardware assets, warranty contracts, and vendors.",
                    "Log maintenance repair invoices and generate cost reports.",
                    "Maintain departments, users, categories, and system priorities."
                ]
            }
        ]
    )

    # SLIDE 8: Functional Requirements
    make_content_slide(
        8, "Software Engineering", "Key Functional Requirements (FR1 - FR7)",
        [
            {
                "header": "FR1 - FR3",
                "title": "Ticketing & Workflow",
                "badge_color": C_BLUE,
                "bullets": [
                    "FR1 (Ticket Creation): Users create tickets with unique ticket_no, category, priority, and subtype classification.",
                    "FR2 (Technician Assignment): Administrators assign pending tickets to specialized support staff.",
                    "FR3 (Status Lifecycle): System logs every state change into status_histories with old/new states."
                ]
            },
            {
                "header": "FR4 - FR5",
                "title": "Resolution & Inventory",
                "badge_color": C_EMERALD,
                "bullets": [
                    "FR4 (Ticket Resolution): Technicians record resolution descriptions and timestamps with unique 1:1 enforcement.",
                    "FR5 (Asset Tracking): Hardware assets cataloged with unique asset_tag, serial_no, custodian user, and warranty."
                ]
            },
            {
                "header": "FR6 - FR7",
                "title": "Maintenance & Reporting",
                "badge_color": C_AMBER,
                "bullets": [
                    "FR6 (Maintenance Logging): Record asset repair dates, descriptions, and expenses with non-negative check constraints.",
                    "FR7 (Relational Reporting): Provide multi-table JOINs, workload summaries, and audit trail extractions."
                ]
            }
        ]
    )

    # SLIDE 9: Expected Benefits & Impact
    make_content_slide(
        9, "Value Proposition", "Expected Operational Benefits & Institutional Impact",
        [
            {
                "header": "Operational Efficiency",
                "title": "Streamlined Support Workflows",
                "badge_color": C_BLUE,
                "bullets": [
                    "Faster Issue Turnaround: Eliminates lost tickets and clarifies technician ownership.",
                    "Automated Accountability: Complete audit trails deter bottlenecks and track resolution performance.",
                    "Reduced Duplicate Work: Structured search and categorization prevent duplicate ticket filings."
                ]
            },
            {
                "header": "Resource Optimization",
                "title": "Hardware & Financial Control",
                "badge_color": C_EMERALD,
                "bullets": [
                    "Zero Lost Equipment: Accurate asset-to-user mappings prevent misplaced or unreturned laptops.",
                    "Warranty Protection: Automated visibility into vendor warranty dates prevents paying for covered repairs.",
                    "Budget Forecasting: Departmental maintenance reports highlight recurring equipment failures."
                ]
            }
        ]
    )

    # SLIDE 10: Conclusion & Next Steps
    make_content_slide(
        10, "Summary & Milestones", "Review 1 Conclusion & Academic Road Map",
        [
            {
                "header": "Review 1 Accomplishments",
                "title": "Requirements Formulated",
                "badge_color": C_EMERALD,
                "bullets": [
                    "Comprehensive problem analysis and stakeholder identification completed.",
                    "Clear boundaries and functional requirements (FR1 - FR7) established.",
                    "Conceptual relational entities identified for 14 tables in 3rd Normal Form."
                ]
            },
            {
                "header": "Subsequent Evaluation Phases",
                "title": "Presentation-II & III Road Map",
                "badge_color": C_BLUE,
                "bullets": [
                    "Presentation-II Deliverables: Complete MySQL DDL schema, 55 relational seed records, ER diagram, and analytical queries.",
                    "Presentation-III Deliverables: Live working user interface connected to MySQL demonstrating real-time INSERT, DELETE, and VIEW operations.",
                    "Final Project Report: Comprehensive 16-chapter academic documentation."
                ]
            }
        ]
    )

    out_path = "Presentation-I/Presentation-I.pptx"
    prs.save(out_path)
    print(f"Saved {out_path}")

build_pptx()
