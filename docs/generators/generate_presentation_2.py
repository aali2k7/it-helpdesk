import os
import subprocess
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def build_pptx():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Palette
    C_BG = RGBColor(248, 250, 252)       # #F8FAFC
    C_CARD = RGBColor(255, 255, 255)     # #FFFFFF
    C_BORDER = RGBColor(226, 232, 240)   # #E2E8F0
    C_NAVY = RGBColor(15, 23, 42)        # #0F172A
    C_SLATE = RGBColor(51, 65, 85)       # #334155
    C_MUTED = RGBColor(100, 116, 139)    # #64748B
    C_BLUE = RGBColor(37, 99, 235)       # #2563EB
    C_BLUE_BG = RGBColor(239, 246, 255)  # #EFF6FF
    C_EMERALD = RGBColor(5, 150, 105)    # #059669
    C_AMBER = RGBColor(217, 119, 6)      # #D97706

    def add_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = C_BG
        bg.line.fill.background()
        return bg

    def add_header(slide, eyebrow, title, slide_num):
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(1.1))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p0 = tf.paragraphs[0]
        p0.text = eyebrow.upper()
        p0.font.size = Pt(10)
        p0.font.bold = True
        p0.font.color.rgb = C_BLUE

        p1 = tf.add_paragraph()
        p1.text = title
        p1.font.size = Pt(22)
        p1.font.bold = True
        p1.font.color.rgb = C_NAVY

        # Footer
        fb = slide.shapes.add_textbox(Inches(0.8), Inches(7.0), Inches(11.733), Inches(0.4))
        ftf = fb.text_frame
        ftf.word_wrap = True
        ftf.margin_left = ftf.margin_top = ftf.margin_right = ftf.margin_bottom = 0
        fp = ftf.paragraphs[0]
        fp.text = f"IT Helpdesk & Asset Support Management System | Presentation-II: Schema Implementation & SQL Queries | Slide {slide_num} of 14"
        fp.font.size = Pt(9)
        fp.font.color.rgb = C_MUTED

    def make_content_slide(slide_num, eyebrow, title, cards_data):
        slide = prs.slides.add_slide(blank_layout)
        add_bg(slide)
        add_header(slide, eyebrow, title, slide_num)

        num_cards = len(cards_data)
        card_w = (11.733 - (num_cards - 1) * 0.3) / num_cards
        for idx, c in enumerate(cards_data):
            cx = 0.8 + idx * (card_w + 0.3)
            box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(cx), Inches(1.6), Inches(card_w), Inches(5.1))
            box.fill.solid()
            box.fill.fore_color.rgb = C_CARD
            box.line.color.rgb = C_BORDER

            btf = box.text_frame
            btf.word_wrap = True
            btf.margin_left = Inches(0.25)
            btf.margin_right = Inches(0.25)
            btf.margin_top = Inches(0.25)

            p = btf.paragraphs[0]
            p.text = c.get('header', '').upper()
            p.font.size = Pt(10)
            p.font.bold = True
            p.font.color.rgb = c.get('badge_color', C_BLUE)

            p = btf.add_paragraph()
            p.text = c.get('title', '')
            p.font.size = Pt(16)
            p.font.bold = True
            p.font.color.rgb = C_NAVY
            p.space_after = Pt(10)

            for bullet in c.get('bullets', []):
                p = btf.add_paragraph()
                p.text = f"• {bullet}"
                p.font.size = Pt(12)
                p.font.color.rgb = C_SLATE
                p.space_after = Pt(6)

        return slide

    # -------------------------------------------------------------
    # SLIDE 1: Title Slide
    # -------------------------------------------------------------
    s1 = prs.slides.add_slide(blank_layout)
    add_bg(s1)

    card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(0.8), Inches(11.333), Inches(5.9))
    card.fill.solid()
    card.fill.fore_color.rgb = C_CARD
    card.line.color.rgb = C_BORDER

    tb = s1.shapes.add_textbox(Inches(1.5), Inches(1.2), Inches(10.333), Inches(3.0))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "DATABASE MANAGEMENT SYSTEMS (DBMS) COURSE PROJECT"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = C_BLUE

    p = tf.add_paragraph()
    p.text = "IT Helpdesk & Asset Support\nManagement System"
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = C_NAVY
    p.space_after = Pt(14)

    p = tf.add_paragraph()
    p.text = "Presentation-II: Relational Schema Design, 3NF Normalization, ER Diagram & SQL Query Implementation"
    p.font.size = Pt(14)
    p.font.color.rgb = C_SLATE

    # Left: Student
    ib1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(4.5), Inches(4.8), Inches(1.8))
    ib1.fill.solid()
    ib1.fill.fore_color.rgb = C_BLUE_BG
    ib1.line.color.rgb = C_BORDER
    itf1 = ib1.text_frame
    itf1.margin_left = Inches(0.2)
    itf1.margin_top = Inches(0.15)
    ip = itf1.paragraphs[0]
    ip.text = "STUDENT DETAILS"
    ip.font.size = Pt(10)
    ip.font.bold = True
    ip.font.color.rgb = C_BLUE

    ip = itf1.add_paragraph()
    ip.text = "Md Aali Rahman"
    ip.font.size = Pt(16)
    ip.font.bold = True
    ip.font.color.rgb = C_NAVY

    ip = itf1.add_paragraph()
    ip.text = "Roll Number: 25WU0102156\nSection: AIML Panthers | Year: 2026\nWoxsen University, School of Technology"
    ip.font.size = Pt(11)
    ip.font.color.rgb = C_SLATE

    # Right: Faculty
    ib2 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(4.5), Inches(5.0), Inches(1.8))
    ib2.fill.solid()
    ib2.fill.fore_color.rgb = C_CARD
    ib2.line.color.rgb = C_BORDER
    itf2 = ib2.text_frame
    itf2.margin_left = Inches(0.2)
    itf2.margin_top = Inches(0.15)
    ip = itf2.paragraphs[0]
    ip.text = "ACADEMIC MENTORSHIP"
    ip.font.size = Pt(10)
    ip.font.bold = True
    ip.font.color.rgb = C_BLUE

    ip = itf2.add_paragraph()
    ip.text = "Faculty Guide: Dr. Kiranmayee Adavala"
    ip.font.size = Pt(15)
    ip.font.bold = True
    ip.font.color.rgb = C_NAVY

    ip = itf2.add_paragraph()
    ip.text = "Department: Computer Science & Engineering\nDatabase Engine: MySQL 8.0+ / 9.x (InnoDB)\nEvaluated Deliverables: DDL, Seed DML, ERD, Queries"
    ip.font.size = Pt(11)
    ip.font.color.rgb = C_SLATE

    # SLIDE 2: Problem / System Overview
    make_content_slide(
        2, "System Architecture", "Problem Recap & Relational Database Scope",
        [
            {
                "header": "Core Problem",
                "title": "IT Service Coordination",
                "badge_color": C_BLUE,
                "bullets": [
                    "Manual, email-driven support causes lost tickets, delayed resolutions, and missing asset records.",
                    "No relational correlation between hardware equipment, warranties, and cumulative maintenance costs.",
                    "Need for a robust relational back-end to enforce data consistency and historical audit trails."
                ]
            },
            {
                "header": "Database Scope",
                "title": "Centralized it_helpdesk Schema",
                "badge_color": C_EMERALD,
                "bullets": [
                    "Centralized MySQL relational database managing 14 distinct tables in 3rd Normal Form (3NF).",
                    "Dual-path ticket classification: Incidents (faults) and Service Requests (access & allocations).",
                    "Full lifecycle tracking with assignments, status histories, and resolution logging."
                ]
            }
        ]
    )

    # SLIDE 3: Database Objectives
    make_content_slide(
        3, "Engineering Goals", "Relational Database Design Objectives",
        [
            {
                "header": "Integrity & Reliability",
                "title": "Data Governance Objectives",
                "badge_color": C_BLUE,
                "bullets": [
                    "Eliminate Data Redundancy: Apply 1NF, 2NF, and 3NF across all entities to achieve single-source data storage.",
                    "Enforce Referential Integrity: Implement foreign key actions (RESTRICT, CASCADE, SET NULL) matching real-world lifecycle semantics.",
                    "Domain Constraints: Enforce CHECK constraints on priority levels (1-5), status enums, and non-negative maintenance costs."
                ]
            },
            {
                "header": "Performance & Auditing",
                "title": "Analytical & Operational Goals",
                "badge_color": C_EMERALD,
                "bullets": [
                    "Audit Trail Accountability: Automatically capture old vs new states with timestamps in status_histories.",
                    "High-Performance Queries: Support multi-table JOINs and GROUP BY aggregation for technician workloads and expenditure.",
                    "Production Readiness: Ensure InnoDB ACID compliance and crash resilience."
                ]
            }
        ]
    )

    # SLIDE 4: ER Diagram Slide (With Image)
    s4 = prs.slides.add_slide(blank_layout)
    add_bg(s4)
    add_header(s4, "Conceptual Model", "Entity-Relationship (ER) Diagram Overview", 4)
    # Card container
    c_box = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(11.733), Inches(5.1))
    c_box.fill.solid()
    c_box.fill.fore_color.rgb = C_CARD
    c_box.line.color.rgb = C_BORDER

    # Insert ER diagram image if exists
    erd_img = "Presentation-II/ER-Diagram.png"
    if os.path.exists(erd_img):
        s4.shapes.add_picture(erd_img, Inches(1.1), Inches(1.8), Inches(8.0), Inches(4.7))
        # Side summary card
        sb = s4.shapes.add_textbox(Inches(9.3), Inches(1.8), Inches(3.0), Inches(4.7))
        sbtf = sb.text_frame
        sbtf.word_wrap = True
        sp = sbtf.paragraphs[0]
        sp.text = "KEY RELATIONSHIPS"
        sp.font.size = Pt(11)
        sp.font.bold = True
        sp.font.color.rgb = C_BLUE
        
        rel_notes = [
            "14 Entities in 3NF",
            "14 Foreign Key relations",
            "departments 1:N users",
            "users 1:N tickets",
            "tickets 1:1 incidents",
            "tickets 1:1 service_requests",
            "tickets 1:1 resolutions",
            "tickets 1:N assignments",
            "tickets 1:N status_histories",
            "assets 1:N maintenance"
        ]
        for rn in rel_notes:
            sp = sbtf.add_paragraph()
            sp.text = f"• {rn}"
            sp.font.size = Pt(11)
            sp.font.color.rgb = C_SLATE

    # SLIDE 5: Relational Schema
    make_content_slide(
        5, "Physical Schema", "Relational Schema Architecture (14 Tables)",
        [
            {
                "header": "Core Master Tables",
                "title": "Master & Reference Tables",
                "badge_color": C_BLUE,
                "bullets": [
                    "departments(department_id PK, name UNIQUE, location)",
                    "users(user_id PK, name, email UNIQUE, department_id FK)",
                    "categories(category_id PK, name UNIQUE, description)",
                    "priorities(priority_id PK, name UNIQUE, level CHECK)",
                    "support_staff(staff_id PK, name, email UNIQUE, specialization)",
                    "warranties(warranty_id PK, start_date, end_date, provider)"
                ]
            },
            {
                "header": "Operational & Transactional",
                "title": "Transactional & Audit Tables",
                "badge_color": C_EMERALD,
                "bullets": [
                    "assets(asset_id PK, asset_tag, serial_no, user_id FK, category_id FK, warranty_id FK)",
                    "tickets(ticket_id PK, ticket_no UNIQUE, user_id FK, category_id FK, priority_id FK, status)",
                    "incidents(ticket_id PK/FK, incident_type)",
                    "service_requests(ticket_id PK/FK, request_type)",
                    "assignments(assignment_id PK, ticket_id FK, staff_id FK, assigned_at)",
                    "status_histories(history_id PK, ticket_id FK, old_status, new_status, changed_at)",
                    "resolutions(resolution_id PK, ticket_id FK UNIQUE, description, resolved_at)",
                    "maintenance(maintenance_id PK, asset_id FK, maintenance_date, cost, description)"
                ]
            }
        ]
    )

    # SLIDE 6: Tables & Relationships
    make_content_slide(
        6, "Referential Integrity", "Foreign Key Semantics & Cascading Rules",
        [
            {
                "header": "RESTRICT Actions",
                "title": "Master Data Protection",
                "badge_color": C_BLUE,
                "bullets": [
                    "departments -> users (ON DELETE RESTRICT): Prevents deleting a department while active employees exist.",
                    "categories -> tickets / assets (ON DELETE RESTRICT): Protects operational categories from accidental deletion.",
                    "priorities -> tickets (ON DELETE RESTRICT): Guarantees priority integrity for active queues.",
                    "support_staff -> assignments (ON DELETE RESTRICT): Preserves technician service records."
                ]
            },
            {
                "header": "CASCADE & SET NULL",
                "title": "Lifecycle Synchronized Rules",
                "badge_color": C_EMERALD,
                "bullets": [
                    "tickets -> incidents / service_requests / assignments / status_histories / resolutions (ON DELETE CASCADE): Deleting a ticket cleanly purges sub-entities without leaving orphan rows.",
                    "assets -> maintenance (ON DELETE CASCADE): Removing an asset purges historical maintenance receipts.",
                    "users -> assets (ON DELETE SET NULL): Removing an employee safely unassigns hardware back into general inventory."
                ]
            }
        ]
    )

    # SLIDE 7: Normalization (1NF, 2NF, 3NF)
    make_content_slide(
        7, "Database Theory", "Normalization Analysis (1NF to 3NF)",
        [
            {
                "header": "1NF & 2NF",
                "title": "First & Second Normal Form",
                "badge_color": C_BLUE,
                "bullets": [
                    "1NF (First Normal Form): All table attributes are atomic. No multi-valued attributes, comma-separated lists, or repeating groups. Every column has defined scalar data types.",
                    "2NF (Second Normal Form): Satisfies 1NF, and all non-key attributes are fully functionally dependent on the entire primary key. Composite keys eliminated in favor of surrogate integer PKs, preventing partial functional dependencies."
                ]
            },
            {
                "header": "3NF",
                "title": "Third Normal Form Compliance",
                "badge_color": C_EMERALD,
                "bullets": [
                    "3NF (Third Normal Form): Satisfies 2NF and eliminates all transitive functional dependencies (X -> Y and Y -> Z).",
                    "Departments are decoupled from users (users store only department_id, not department location).",
                    "Warranties are isolated into a separate entity (assets store warranty_id rather than repeating provider and dates).",
                    "Categories and priorities are independent reference tables."
                ]
            }
        ]
    )

    # SLIDE 8: DDL Commands
    make_content_slide(
        8, "Implementation", "Data Definition Language (DDL) Architecture",
        [
            {
                "header": "Schema Construction",
                "title": "Table Creation & Data Types",
                "badge_color": C_BLUE,
                "bullets": [
                    "Engineered with MySQL InnoDB storage engine for transactional safety and row-level locking.",
                    "Character set configured to utf8mb4 with utf8mb4_unicode_ci collation for universal Unicode text support.",
                    "AUTO_INCREMENT surrogate primary keys for high-efficiency B-Tree indexing.",
                    "Appropriate sizing: VARCHAR(100), VARCHAR(150), TEXT for narratives, TIMESTAMP for audit precision."
                ]
            },
            {
                "header": "Constraints Enforcement",
                "title": "DDL Constraints Verification",
                "badge_color": C_AMBER,
                "bullets": [
                    "CHECK (level BETWEEN 1 AND 5) on table priorities.",
                    "CHECK (status IN ('Open', 'In Progress', 'Pending', 'Resolved', 'Closed', 'Cancelled')) on table tickets.",
                    "CHECK (end_date >= start_date) on table warranties.",
                    "CHECK (cost >= 0) on table maintenance.",
                    "UNIQUE constraints on emails, ticket numbers, asset tags, and serial numbers."
                ]
            }
        ]
    )

    # SLIDE 9: DML / INSERT Commands
    make_content_slide(
        9, "Data Manipulation", "DML Seed Dataset Implementation",
        [
            {
                "header": "Seed Dataset",
                "title": "55 Relational Seed Records",
                "badge_color": C_BLUE,
                "bullets": [
                    "Structured seed dataset in seed.sql establishing realistic operational baseline across 14 tables.",
                    "4 Departments: Engineering, Human Resources, Finance, Marketing.",
                    "5 Registered Corporate Users across all departments.",
                    "5 Ticket Categories & 4 Priority Levels (Low, Medium, High, Critical).",
                    "4 Support Staff Technicians with distinct specializations (Network, Hardware, Software, IAM)."
                ]
            },
            {
                "header": "Referential Linking",
                "title": "Integrated Relational Records",
                "badge_color": C_EMERALD,
                "bullets": [
                    "4 Manufacturer Warranties (Dell ProSupport, AppleCare, Lenovo Premier, HP Care Pack).",
                    "5 Computing Assets linked to user custodians and warranty providers.",
                    "5 Base Tickets spanning Incidents and Service Requests with initial status histories.",
                    "Active technician assignments and recorded maintenance repairs."
                ]
            }
        ]
    )

    # SLIDE 10: SELECT Queries
    make_content_slide(
        10, "Query Optimization", "Representative SELECT & Multi-Table JOIN Queries",
        [
            {
                "header": "Relational JOINs",
                "title": "Multi-Table Ticket Queue Query",
                "badge_color": C_BLUE,
                "bullets": [
                    "Joins tickets, users, departments, categories, priorities, incidents, and service_requests.",
                    "Utilizes CASE expressions to determine ticket subtype dynamically (Incident vs Service Request).",
                    "Extracts human-readable status queues ordered by priority level descending.",
                    "Preserves NULL values using LEFT OUTER JOINs for unassigned subtypes."
                ]
            },
            {
                "header": "Aggregate Analytics",
                "title": "Workload & Expenditure Aggregation",
                "badge_color": C_EMERALD,
                "bullets": [
                    "Technician Workload Query: Uses GROUP BY with COUNT and conditional SUM(CASE...) to calculate active vs resolved tickets per engineer.",
                    "Department Maintenance Spend: Aggregates total equipment repair cost per department using SUM(m.cost) and FORMAT().",
                    "Status Distribution: Calculates real-time percentage distribution across ticket lifecycle states."
                ]
            }
        ]
    )

    # SLIDE 11: Presentation-II Assigned Query
    make_content_slide(
        11, "Evaluation Review", "Presentation-II Assigned Viva Query",
        [
            {
                "header": "Evaluation Status",
                "title": "Assigned Query Notice",
                "badge_color": C_AMBER,
                "bullets": [
                    "During Presentation-II evaluation, each learner was assigned a specific SQL problem statement to solve live.",
                    "The official template in Presentation-II/Presentation-II-Query-Solution.sql archives the query, explanation, and output.",
                    "Repository status: Ready for student viva query parameters as reviewed by faculty.",
                    "System provides the live SQL Execution Playground to test and demonstrate any complex query."
                ]
            },
            {
                "header": "Core Problem Pattern",
                "title": "Analytical Helpdesk Requirements",
                "badge_color": C_BLUE,
                "bullets": [
                    "Focuses on multi-table joins combining operational tickets with technician workloads.",
                    "Filtering on date ranges, ticket statuses, or departmental asset allocations.",
                    "Demonstrates aggregate grouping (GROUP BY) and conditional aggregation (HAVING)."
                ]
            }
        ]
    )

    # SLIDE 12: Query Output & Verification
    make_content_slide(
        12, "Query Execution", "Query Output Verification & Performance",
        [
            {
                "header": "Live Execution",
                "title": "Verified Database Output",
                "badge_color": C_EMERALD,
                "bullets": [
                    "Technician Workload Summary Output:",
                    "• Sneha Rao (Network & Security): 1 Total, 1 Active, 0 Completed",
                    "• Karan Malhotra (Hardware & Peripherals): 2 Total, 2 Active, 0 Completed",
                    "• Amit Joshi (OS & Software): 1 Total, 1 Active, 0 Completed",
                    "• Divya Nair (IAM): 1 Total, 0 Active, 1 Completed"
                ]
            },
            {
                "header": "Cost Output",
                "title": "Departmental Maintenance Spend",
                "badge_color": C_BLUE,
                "bullets": [
                    "• Finance & Accounts: 1 Asset | $1,200.00 Maintenance Spent",
                    "• Engineering: 2 Assets | $750.00 Maintenance Spent",
                    "• Marketing & Sales: 1 Asset | $450.00 Maintenance Spent",
                    "• Human Resources: 1 Asset | $0.00 Maintenance Spent"
                ]
            }
        ]
    )

    # SLIDE 13: Database Summary
    make_content_slide(
        13, "System Summary", "Database Architecture Metrics & Implementation Summary",
        [
            {
                "header": "Schema Metrics",
                "title": "14 Normalized Tables",
                "badge_color": C_BLUE,
                "bullets": [
                    "Tables: departments, users, categories, priorities, support_staff, warranties, assets, tickets, incidents, service_requests, assignments, status_histories, resolutions, maintenance.",
                    "Relationships: 14 Foreign Keys enforcing strict referential integrity.",
                    "Constraints: 4 Domain CHECK constraints, 8 UNIQUE constraints, 14 Primary Keys."
                ]
            },
            {
                "header": "Storage & Engine",
                "title": "MySQL InnoDB Implementation",
                "badge_color": C_EMERALD,
                "bullets": [
                    "55 Relational seed records ensuring immediate operational verification.",
                    "Terminal CLI interface implemented in main.py (--cli mode) with full tabular views.",
                    "Live MySQL connectivity verified on port 3306."
                ]
            }
        ]
    )

    # SLIDE 14: Conclusion & Next Steps
    make_content_slide(
        14, "Conclusion", "Presentation-II Conclusion & Transition to UI Demo",
        [
            {
                "header": "Milestones Achieved",
                "title": "Presentation-II Deliverables Finalized",
                "badge_color": C_EMERALD,
                "bullets": [
                    "Robust, 3NF-compliant relational database fully implemented in MySQL.",
                    "Complete DDL schema (schema.sql) and seed dataset (seed.sql) operational.",
                    "Interactive ER diagram and high-resolution exports generated.",
                    "Analytical SQL queries tested and documented."
                ]
            },
            {
                "header": "Next Phase",
                "title": "Presentation-III (UI Demonstration)",
                "badge_color": C_BLUE,
                "bullets": [
                    "Development of interactive Web User Interface connected to MySQL.",
                    "Live demonstration of INSERT, DELETE, and VIEW operations with real-time database reflection.",
                    "Preparation of complete 16-chapter Project Report PDF."
                ]
            }
        ]
    )

    out_path = "Presentation-II/Presentation-II.pptx"
    prs.save(out_path)
    print(f"Saved {out_path}")

build_pptx()
